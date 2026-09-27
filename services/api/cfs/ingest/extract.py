"""Deterministic content extraction: file-type detection and page-level text."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass

from cfs.ingest.text import normalize

PARSER_VERSION = "extract-1.0"
MIN_CHARS_PER_PAGE = 40  # below this a PDF page is treated as image-only / low-quality


class UnsupportedFile(Exception):
    pass


@dataclass
class PageText:
    page_number: int
    text: str
    method: str
    width: float | None = None
    height: float | None = None
    low_quality: bool = False


def detect_mime(data: bytes, filename: str | None = None) -> str:
    """Sniff content; never trust the client-supplied MIME type or extension alone."""
    if data[:5] == b"%PDF-":
        return "application/pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except zipfile.BadZipFile:
            pass
        raise UnsupportedFile("ZIP container is not a DOCX document")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as e:
        raise UnsupportedFile("Unsupported file type (expected PDF, DOCX, or UTF-8 text)") from e
    if "\x00" in text:
        raise UnsupportedFile("Binary content is not supported")
    return "text/markdown" if (filename or "").lower().endswith(".md") else "text/plain"


EXT_FOR_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
}


def extract_pages(data: bytes, mime: str) -> list[PageText]:
    if mime == "application/pdf":
        return _pdf_pages(data)
    if mime.endswith("wordprocessingml.document"):
        return _docx_pages(data)
    return _text_pages(data.decode("utf-8"), "text")


def _pdf_pages(data: bytes) -> list[PageText]:
    import pdfplumber

    out: list[PageText] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            low = len(normalize(text)) < MIN_CHARS_PER_PAGE
            out.append(PageText(i, text, "pdfplumber", float(page.width), float(page.height), low))
    return out


def _docx_pages(data: bytes) -> list[PageText]:
    import docx

    d = docx.Document(io.BytesIO(data))
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            paras.append(" | ".join(c.text for c in row.cells))
    return _text_pages("\n".join(paras), "python-docx")


def _text_pages(text: str, method: str, page_chars: int = 3000) -> list[PageText]:
    """DOCX/text have no physical pages: split on form feeds, else into ~3000-char sections at line breaks."""
    if "\f" in text:
        chunks = text.split("\f")
    else:
        chunks, cur = [], ""
        for line in text.splitlines(keepends=True):
            if len(cur) + len(line) > page_chars and cur:
                chunks.append(cur)
                cur = ""
            cur += line
        chunks.append(cur)
    return [PageText(i, c, method) for i, c in enumerate(chunks, start=1) if c.strip() or i == 1]


def pdf_line_boxes(data: bytes, page_number: int, quote: str) -> list[list[float]] | None:
    """Bounding boxes (PDF points, top-left origin) of each line of a quote on a page, for highlighting."""
    import re

    import pdfplumber

    from cfs.ingest.text import _flex_pattern

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if page_number < 1 or page_number > len(pdf.pages):
            return None
        page = pdf.pages[page_number - 1]
        pattern = _flex_pattern(quote).pattern
        try:
            hits = page.search(pattern, regex=True, case=False, return_chars=True)
        except re.error:
            return None
        if not hits:
            return None
        chars = hits[0]["chars"]
        lines: dict[int, list[float]] = {}
        for c in chars:
            key = round(c["top"])
            b = lines.setdefault(key, [c["x0"], c["top"], c["x1"], c["bottom"]])
            b[0], b[1] = min(b[0], c["x0"]), min(b[1], c["top"])
            b[2], b[3] = max(b[2], c["x1"]), max(b[3], c["bottom"])
        return [[round(v, 2) for v in b] for b in lines.values()]
